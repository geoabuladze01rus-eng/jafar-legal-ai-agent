from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path


@dataclass(frozen=True, slots=True)
class DocumentFragment:
    text: str
    chunk_index: int
    page: int | None = None


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    media_type: str
    text: str
    pages: tuple[str, ...] = ()

    @property
    def fingerprint(self) -> str:
        """Stable identity for the extracted document content."""
        payload = f"{self.media_type}\n{self.text}".encode()
        return sha256(payload).hexdigest()

    def fragments(self, *, max_chars: int = 2_000) -> tuple[DocumentFragment, ...]:
        """Return stable page-aware chunks without changing the legacy full-text contract."""
        if max_chars < 200:
            raise ValueError("max_chars must be at least 200")
        units = self.pages or (self.text,)
        result: list[DocumentFragment] = []
        chunk_index = 0
        for page_index, unit in enumerate(units, start=1):
            normalized = "\n".join(line.rstrip() for line in unit.splitlines()).strip()
            if not normalized:
                continue
            start = 0
            while start < len(normalized):
                end = min(len(normalized), start + max_chars)
                if end < len(normalized):
                    break_at = normalized.rfind("\n", start, end)
                    if break_at <= start:
                        break_at = normalized.rfind(" ", start, end)
                    if break_at > start + max_chars // 2:
                        end = break_at
                fragment_text = normalized[start:end].strip()
                if fragment_text:
                    result.append(
                        DocumentFragment(
                            text=fragment_text,
                            chunk_index=chunk_index,
                            page=page_index if self.pages else None,
                        )
                    )
                    chunk_index += 1
                start = max(end, start + 1)
        return tuple(result)

    def evidence_id(self, fragment: DocumentFragment) -> str:
        page_part = f":page:{fragment.page}" if fragment.page is not None else ""
        return f"document:{self.fingerprint}{page_part}:chunk:{fragment.chunk_index}"


class DocumentExtractionError(ValueError):
    pass


class DocumentExtractor:
    """Extract text from legal documents supported by Jafar."""

    MAX_BYTES = 20 * 1024 * 1024
    SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx"}

    def extract(self, filename: str, content: bytes, media_type: str | None = None) -> ExtractedDocument:
        if len(content) > self.MAX_BYTES:
            raise DocumentExtractionError("Document exceeds the 20 MB limit")
        extension = Path(filename).suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise DocumentExtractionError(f"Unsupported document format: {extension or 'unknown'}")
        pages: tuple[str, ...] = ()
        if extension in {".txt", ".md", ".markdown"}:
            text = content.decode("utf-8-sig", errors="replace")
        elif extension == ".pdf":
            pages = self._extract_pdf_pages(content)
            text = "\n".join(pages)
        else:
            text = self._extract_docx(content)
        text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
        pages = tuple(
            "\n".join(line.rstrip() for line in page.splitlines()).strip() for page in pages
        )
        if not text:
            raise DocumentExtractionError("No text could be extracted from the document")
        return ExtractedDocument(
            filename=filename,
            media_type=media_type or "application/octet-stream",
            text=text,
            pages=pages,
        )

    @staticmethod
    def _extract_pdf_pages(content: bytes) -> tuple[str, ...]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentExtractionError("PDF support is not installed") from exc
        try:
            reader = PdfReader(BytesIO(content))
            return tuple(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise DocumentExtractionError("Unable to read PDF document") from exc

    @classmethod
    def _extract_pdf(cls, content: bytes) -> str:
        """Compatibility helper retained for callers that expect full PDF text."""
        return "\n".join(cls._extract_pdf_pages(content))

    @staticmethod
    def _extract_docx(content: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentExtractionError("DOCX support is not installed") from exc
        try:
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(paragraphs)
        except Exception as exc:
            raise DocumentExtractionError("Unable to read DOCX document") from exc
