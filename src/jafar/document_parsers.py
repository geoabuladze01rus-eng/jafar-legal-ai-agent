from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Protocol


class DocumentParserError(ValueError):
    pass


class DocumentParser(Protocol):
    def supports(self, filename: str, media_type: str | None = None) -> bool: ...

    def parse(self, filename: str, content: bytes, media_type: str | None = None) -> str: ...


@dataclass(slots=True)
class NativeDocumentParser:
    """Small, dependency-light parser used by the existing Jafar pipeline."""

    def supports(self, filename: str, media_type: str | None = None) -> bool:
        return Path(filename).suffix.lower() in {".txt", ".md", ".markdown", ".pdf", ".docx"}

    def parse(self, filename: str, content: bytes, media_type: str | None = None) -> str:
        extension = Path(filename).suffix.lower()
        if extension in {".txt", ".md", ".markdown"}:
            return content.decode("utf-8-sig", errors="replace")
        if extension == ".pdf":
            return self._extract_pdf(content)
        if extension == ".docx":
            return self._extract_docx(content)
        raise DocumentParserError(f"Unsupported document format: {extension or 'unknown'}")

    @staticmethod
    def _extract_pdf(content: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentParserError("PDF support is not installed") from exc
        try:
            reader = PdfReader(BytesIO(content))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise DocumentParserError("Unable to read PDF document") from exc

    @staticmethod
    def _extract_docx(content: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentParserError("DOCX support is not installed") from exc
        try:
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(paragraphs)
        except Exception as exc:
            raise DocumentParserError("Unable to read DOCX document") from exc


class DoclingDocumentParser:
    """Local-first rich parser for layout-aware legal-document ingestion.

    Docling is imported lazily so the base Jafar install and CI remain lightweight.
    Install it with ``pip install -e '.[docling]'`` when local rich parsing is enabled.
    """

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm"}

    def __init__(self, converter=None, stream_factory=None) -> None:
        self._converter = converter
        self._stream_factory = stream_factory

    def supports(self, filename: str, media_type: str | None = None) -> bool:
        return Path(filename).suffix.lower() in self.SUPPORTED_EXTENSIONS

    def parse(self, filename: str, content: bytes, media_type: str | None = None) -> str:
        converter = self._converter or self._build_converter()
        stream_factory = self._stream_factory or self._build_stream_factory()
        try:
            source = stream_factory(name=filename, stream=BytesIO(content))
            result = converter.convert(source, max_file_size=len(content))
            document = result.document
            if hasattr(document, "export_to_markdown"):
                return document.export_to_markdown()
            if hasattr(document, "export_to_text"):
                return document.export_to_text()
            raise DocumentParserError("Docling returned a document without text export")
        except DocumentParserError:
            raise
        except Exception as exc:
            raise DocumentParserError("Docling could not parse the document") from exc

    @staticmethod
    def _build_converter():
        try:
            from docling.document_converter import DocumentConverter
        except ImportError as exc:
            raise DocumentParserError(
                "Docling support is not installed; install the 'docling' optional dependency"
            ) from exc
        return DocumentConverter()

    @staticmethod
    def _build_stream_factory():
        try:
            from docling.datamodel.base_models import DocumentStream
        except ImportError as exc:
            raise DocumentParserError(
                "Docling support is not installed; install the 'docling' optional dependency"
            ) from exc
        return DocumentStream


@dataclass(slots=True)
class FallbackDocumentParser:
    """Try a primary parser first and preserve the existing native parser as fallback."""

    primary: DocumentParser
    fallback: DocumentParser

    def supports(self, filename: str, media_type: str | None = None) -> bool:
        return self.primary.supports(filename, media_type) or self.fallback.supports(filename, media_type)

    def parse(self, filename: str, content: bytes, media_type: str | None = None) -> str:
        if self.primary.supports(filename, media_type):
            try:
                return self.primary.parse(filename, content, media_type)
            except DocumentParserError:
                if not self.fallback.supports(filename, media_type):
                    raise
        return self.fallback.parse(filename, content, media_type)


def build_document_parser(*, prefer_docling: bool = False) -> DocumentParser:
    native = NativeDocumentParser()
    if not prefer_docling:
        return native
    return FallbackDocumentParser(primary=DoclingDocumentParser(), fallback=native)
