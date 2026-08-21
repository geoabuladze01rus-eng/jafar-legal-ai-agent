from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass(frozen=True)
class SourceSpan:
    """A traceable text segment that can be cited in a legal analysis."""

    source: str
    page: int | None
    text: str


@dataclass(frozen=True)
class IngestedDocument:
    filename: str
    media_type: str
    text: str
    sources: tuple[SourceSpan, ...]
    ocr_required: bool = False


class UnsupportedDocumentError(ValueError):
    pass


def _media_type(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    return {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".txt": "text/plain",
        ".md": "text/markdown",
    }.get(suffix, "application/octet-stream")


def ingest_document(filename: str, content: bytes) -> IngestedDocument:
    """Extract text while preserving page-level provenance where available.

    Scanned PDFs are detected when pages contain no extractable text. OCR is
    deliberately represented as a separate step so a production OCR provider
    can be added without changing the document API.
    """
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentError(f"Unsupported document type: {suffix or 'unknown'}")

    media_type = _media_type(filename)
    if suffix == ".pdf":
        return _ingest_pdf(filename, media_type, content)
    if suffix == ".docx":
        return _ingest_docx(filename, media_type, content)

    text = content.decode("utf-8-sig", errors="replace").strip()
    source = SourceSpan(source=filename, page=None, text=text)
    return IngestedDocument(filename, media_type, text, (source,), not bool(text))


def _ingest_pdf(filename: str, media_type: str, content: bytes) -> IngestedDocument:
    reader = PdfReader(BytesIO(content))
    spans: list[SourceSpan] = []
    pages: list[str] = []
    empty_pages = 0

    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            empty_pages += 1
        else:
            pages.append(text)
            spans.append(SourceSpan(source=filename, page=index, text=text))

    text = "\n\n".join(pages).strip()
    ocr_required = bool(reader.pages) and empty_pages == len(reader.pages)
    return IngestedDocument(filename, media_type, text, tuple(spans), ocr_required)


def _ingest_docx(filename: str, media_type: str, content: bytes) -> IngestedDocument:
    document = Document(BytesIO(content))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    source = SourceSpan(source=filename, page=None, text=text)
    return IngestedDocument(filename, media_type, text, (source,), not bool(text))
